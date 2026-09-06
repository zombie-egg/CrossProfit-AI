from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from backend.app.database import SessionLocal
from backend.app.models import AnalysisResult, PromotionActivity

from backend.app.main import app


def test_health():
    with TestClient(app) as client:
        response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_products_seeded():
    with TestClient(app) as client:
        response = client.get("/products")
    assert response.status_code == 200
    assert len(response.json()) >= 3


def test_modern_ui_bootstrap_and_combined_analysis():
    with TestClient(app) as client:
        bootstrap = client.get("/ui/bootstrap")
        assert bootstrap.status_code == 200
        data = bootstrap.json()
        assert len(data["products"]) >= 3
        assert len(data["cases"]) >= 3
        case = data["cases"][0]
        response = client.post(
            f"/analysis/run?product_id={case['product_id']}&activity_id={case['id']}",
            json=case["request"],
        )
        analysis_id = response.json()["analysis_id"]
        exported = client.get(f"/export/{analysis_id}?format=xlsx")
    assert response.status_code == 200
    result = response.json()
    assert result["result"]["unit_profit"] == case["result"]["unit_profit"]
    assert len(result["scenarios"]) == 3
    assert result["recommendations"]
    assert result["analysis_id"] is not None
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert workbook.sheetnames == ["Summary", "Cost Breakdown", "Scenario Analysis", "Recommendations"]
    # This integration test persists briefly to verify history/export, then cleans up
    # so repeated test runs never distort real Dashboard statistics.
    with SessionLocal() as session:
        record = session.get(AnalysisResult, analysis_id)
        if record:
            session.delete(record)
            session.commit()


def test_preview_does_not_persist_and_archive_updates_history():
    with TestClient(app) as client:
        case = client.get("/ui/bootstrap").json()["cases"][0]
        before = client.get("/analysis").json()

        preview = client.post("/analysis/run", json=case["request"])
        assert preview.status_code == 200
        assert preview.json()["analysis_id"] is None
        assert len(client.get("/analysis").json()) == len(before)

        archived = client.post(
            f"/analysis/archive?product_id={case['product_id']}&activity_id={case['id']}",
            json=case["request"],
        )
        assert archived.status_code == 200
        archived_data = archived.json()
        assert archived_data["analysis_id"] is not None
        assert archived_data["activity_id"] == case["id"]
        updated_history = client.get("/analysis").json()
        assert len(updated_history) == len(before) + 1
        assert updated_history[0]["id"] == archived_data["analysis_id"]
        assert updated_history[0]["cost_drivers"]
        assert all(float(item["amount"]) > 0 for item in updated_history[0]["cost_drivers"])

    with SessionLocal() as session:
        record = session.get(AnalysisResult, archived_data["analysis_id"])
        if record:
            session.delete(record)
            session.commit()


def test_archive_new_activity_is_atomic_and_visible():
    with TestClient(app) as client:
        case = client.get("/ui/bootstrap").json()["cases"][0]
        payload = case["request"]
        payload["activity"] = {**payload["activity"], "activity_name": "归档联调测试活动"}
        archived = client.post(f"/analysis/archive?product_id={case['product_id']}", json=payload)
        assert archived.status_code == 200
        archived_data = archived.json()
        history = client.get("/analysis").json()
        assert any(item["id"] == archived_data["analysis_id"] for item in history)

    with SessionLocal() as session:
        record = session.get(AnalysisResult, archived_data["analysis_id"])
        activity = session.get(PromotionActivity, archived_data["activity_id"])
        if record:
            session.delete(record)
            session.flush()
        if activity:
            session.delete(activity)
        session.commit()
