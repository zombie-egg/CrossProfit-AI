from io import BytesIO
from decimal import Decimal

from openpyxl import load_workbook

from backend.app.services.export_service import ExportService
from backend.app.services.profit_engine import ProfitEngine
from backend.app.services.scenario_engine import ScenarioEngine
from backend.app.services.strategy_engine import StrategyEngine


def test_strategy_is_grounded_and_export_has_four_sheets(base_request):
    result = ProfitEngine().calculate(base_request)
    scenarios = ScenarioEngine().analyze(base_request)
    tips = StrategyEngine().generate(base_request, result, scenarios)
    assert tips
    payload = ExportService().to_xlsx(base_request, result, scenarios, tips)
    workbook = load_workbook(BytesIO(payload), read_only=True)
    assert workbook.sheetnames == ["Summary", "Cost Breakdown", "Scenario Analysis", "Recommendations"]


def test_loss_strategy_mentions_break_even(base_request):
    base_request.activity.discount_value = Decimal("0.5")
    base_request.activity.creator_commission_rate = Decimal("0.3")
    result = ProfitEngine().calculate(base_request)
    tips = StrategyEngine().generate(base_request, result, ScenarioEngine().analyze(base_request))
    assert any("活动售价至少提高" in tip for tip in tips)
